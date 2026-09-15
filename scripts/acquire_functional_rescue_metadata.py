#!/usr/bin/env python3
"""Acquire an immutable public PD intervention snapshot and audit metadata.

This stage does not compute biological effects or execute author notebooks.
The snapshot is for assessing a new analysis, not a confirmed rescue result.
"""
import argparse
import csv
import hashlib
import json
import os
from pathlib import Path
from urllib.request import Request, urlopen

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs/v2/functional_rescue"
META_COLUMNS = {"Plate", "Well", "tags", "SC_name", "SC_Target_name", "final_conc",
                "Compound", "Concentration (M)", "Destination Plate Barcode", "CpdID"}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cache", type=Path, default=ROOT / "data_raw/functional_rescue")
    parser.add_argument("--output", type=Path, default=DOCS / "metadata_inventory.json")
    parser.add_argument("--fetch", action="store_true")
    args = parser.parse_args()
    manifest = json.loads((DOCS / "input_manifest.json").read_text())
    inventory = {"source_commit": manifest["source_commit"],
                 "scope": "CSV headers and design labels; workbook sheet names, dimensions and headers; no effect calculation",
                 "files": {}}
    for item in manifest["files"]:
        relative = Path(item["path"])
        if relative.is_absolute() or ".." in relative.parts:
            raise ValueError("Unsafe manifest path")
        path = args.cache / relative
        if not path.exists():
            if not args.fetch:
                raise FileNotFoundError(f"{relative}: supply --fetch or the audited cache")
            path.parent.mkdir(parents=True, exist_ok=True)
            raw = urlopen(Request(item["url"], headers={"User-Agent": "pd-envtox-metadata-audit"}),
                          timeout=45).read()
            if hashlib.sha256(raw).hexdigest() != item["sha256"]:
                raise ValueError(f"Source checksum mismatch: {relative}")
            temporary = path.with_name(path.name + ".partial")
            temporary.write_bytes(raw)
            os.replace(temporary, path)
        if digest(path) != item["sha256"]:
            raise ValueError(f"Cached checksum mismatch: {relative}")
        record = {"sha256": item["sha256"], "bytes": path.stat().st_size}
        if path.suffix == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as handle:
                first = handle.readline()
                delimiter = ";" if first.count(";") > first.count(",") else ","
                handle.seek(0)
                reader = csv.DictReader(handle, delimiter=delimiter)
                fields = reader.fieldnames
                values = {k: set() for k in fields if k in META_COLUMNS}
                count = 0
                for row in reader:
                    count += 1
                    for k in values:
                        values[k].add(row[k])
                record.update(rows=count, columns=fields, metadata={
                    k: {"distinct": len(v), "values": sorted(v)}
                    for k, v in sorted(values.items()) if k != "Well"
                })
                if "Well" in values:
                    record["distinct_wells"] = len(values["Well"])
        elif path.suffix == ".xlsx":
            wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
            record["sheets"] = []
            for ws in wb:
                # Read the header only; body values are not part of this audit.
                header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()))
                record["sheets"].append({"name": ws.title, "rows_including_header": ws.max_row,
                                         "columns": ws.max_column, "header": header})
            wb.close()
        else:
            record["inspection"] = "checksum only; notebook outputs not used"
        inventory["files"][item["path"]] = record
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(inventory, indent=2, ensure_ascii=False) + "\n")
    print(f"Verified {len(inventory['files'])} source files; metadata inventory written.")


if __name__ == "__main__":
    main()
