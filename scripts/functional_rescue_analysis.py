#!/usr/bin/env python3
"""Run the fixed matched-dose phenotype benchmark and descriptive assay audit."""
import argparse
import gzip
import hashlib
import json
import math
from pathlib import Path
import re

import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import rankdata

from functional_rescue_design_audit import CACHE, DOCS, MAIN, ROOT, load_design, well

CONFIG = json.loads((DOCS / 'analysis_config_v2.json').read_text())
FEATURES = CONFIG['features']
VALUES = FEATURES + CONFIG['viability_features']
THRESHOLDS = [CONFIG['retention_threshold']] + CONFIG['retention_sensitivities']


def write_table(df, path):
    text = df.to_csv(index=False, sep='\t', na_rep='NA', float_format='%.12g', lineterminator='\n')
    raw = text.encode()
    if path.suffix == '.gz':
        packed = gzip.compress(raw, mtime=0)
        assert gzip.decompress(packed) == raw
        path.write_bytes(packed)
    else:
        path.write_bytes(raw)


def normalize(df, mut_tag, wt_tag):
    df = df.copy()
    df['Well'] = df['Well'].map(well)
    for field in VALUES:
        df[field] = pd.to_numeric(df[field], errors='coerce')
    df['valid'] = np.isfinite(df[VALUES]).all(axis=1) & (df['Nuclei_Number_Living'] >= 0)
    audits = []
    for plate, g in df.groupby('Plate', sort=True):
        mut = g[(g.tags == mut_tag) & g.valid]
        wt = g[(g.tags == wt_tag) & g.valid]
        if min(len(mut), len(wt)) < CONFIG['minimum_control_wells_per_plate_per_genotype']:
            raise ValueError(f'Insufficient valid genotype controls: {plate}')
        med = mut[VALUES].median()
        if med['Nuclei_Number_Living'] <= 0:
            raise ValueError(f'Nonpositive mutant count median: {plate}')
        variance = ((mut[VALUES]-mut[VALUES].mean()).pow(2).sum() +
                    (wt[VALUES]-wt[VALUES].mean()).pow(2).sum())/(len(mut)+len(wt)-2)
        scales = np.sqrt(variance)
        if not np.isfinite(scales).all() or not (scales > 0).all():
            raise ValueError(f'Degenerate control feature scale: {plate}')
        for j, field in enumerate(VALUES):
            df.loc[g.index, f'z{j}'] = (g[field]-med[field])/scales[field]
            if j == 4:
                df.loc[g.index, 'ratio4'] = g[field]/med[field]
            audits.append({'Plate':plate, 'feature':field, 'mutant_median':med[field],
                           'corrected_median':wt[field].median(), 'mutant_control_wells':len(mut),
                           'corrected_control_wells':len(wt), 'pooled_control_sd':scales[field],
                           'invalid_wells':int((~g.valid).sum())})
    return df, pd.DataFrame(audits)


def fit_model(control, shrinkage):
    cols = [f'z{i}' for i in range(len(FEATURES))]
    a = control[(control.tags=='Mut - DMSO') & control.valid][cols].to_numpy()
    b = control[(control.tags=='WT - DMSO') & control.valid][cols].to_numpy()
    m0, m1 = a.mean(axis=0), b.mean(axis=0)
    delta = m1 - m0
    if not np.isfinite(delta).all() or delta[0] >= 0:
        raise ValueError('SNCA calibration gate failed: corrected mean must be lower')
    residual = np.vstack([a-m0, b-m1])
    cov = residual.T @ residual / (len(a)+len(b)-2)
    reg = (1-shrinkage)*cov + shrinkage*np.diag(np.diag(cov)) + 1e-8*np.eye(len(FEATURES))
    w = np.linalg.solve(reg, delta)
    den = float(delta @ w)
    if not np.isfinite(den) or den <= 1e-12:
        raise ValueError('Degenerate phenotype discriminant')
    return {'features':FEATURES, 'shrinkage':shrinkage, 'mutant_mean':m0.tolist(),
            'corrected_mean':m1.tolist(), 'weights':(w/den).tolist(),
            'pooled_covariance':cov.tolist(), 'training_mutant_wells':len(a),
            'training_corrected_wells':len(b)}


def score(df, model):
    df = df.copy()
    x = df[[f'z{i}' for i in range(len(FEATURES))]].to_numpy()
    m0, m1 = np.array(model['mutant_mean']), np.array(model['corrected_mean'])
    df['alpha_score'] = (x[:,0]-m0[0]) / (m1[0]-m0[0])
    df['multivariate_score'] = (x-m0) @ np.array(model['weights'])
    df.loc[~df.valid, ['alpha_score','multivariate_score']] = np.nan
    df['alpha_response'] = df.valid & (df.alpha_score >= CONFIG['alpha_rescue_threshold'])
    for threshold in THRESHOLDS:
        tag = str(threshold)
        tol = CONFIG['integrity_tolerance_sd'][tag]
        df['integrity_'+tag] = df.valid & (df[[f'z{i}' for i in (1,2,3)]] >= -tol).all(axis=1)
        df['viability_'+tag] = df.valid & (df.ratio4 >= threshold) & (df.z5 >= -tol)
        df['joint_'+tag] = df.alpha_response & df['integrity_'+tag] & df['viability_'+tag]
    return df


def aggregate(df):
    rows = []
    for (cat, dose), g in df[df.SC_cat.notna()].groupby(['SC_cat','dose_M'], sort=True):
        valid = g[g.valid]
        row = {'SC_cat':cat,'dose_M':dose,'compound_name':g.compound_name.iloc[0],
               'available_wells':len(g),'valid_wells':len(valid),'plates':','.join(sorted(g.Plate.unique())),
               'eligible':len(valid)>=CONFIG['minimum_compound_wells'],
               'alpha_score':valid.alpha_score.median(), 'multivariate_score':valid.multivariate_score.median(),
               'alpha_fraction':valid.alpha_response.mean()}
        row['alpha_response'] = row['eligible'] and row['alpha_fraction'] >= CONFIG['required_well_fraction']
        for t in THRESHOLDS:
            tag = str(t)
            for kind in ['integrity','viability','joint']:
                row[f'{kind}_fraction_{tag}'] = valid[f'{kind}_{tag}'].mean()
            row[f'joint_response_{tag}'] = row['eligible'] and row[f'joint_fraction_{tag}'] >= CONFIG['required_well_fraction']
        for i, name in enumerate(VALUES):
            row[name+'_standardized'] = valid[f'z{i}'].median()
        row['Nuclei_Number_Living_ratio'] = valid.ratio4.median()
        rows.append(row)
    return pd.DataFrame(rows)


def average_precision(y, s):
    y, s = np.asarray(y, dtype=bool), np.asarray(s, dtype=float)
    if not y.sum():
        return float('nan')
    order = np.argsort(-s, kind='stable')
    yy, ss = y[order], s[order]
    # Evaluate at the end of each tied score group.
    ends = np.r_[np.flatnonzero(ss[:-1] != ss[1:]), len(ss)-1]
    tp = np.cumsum(yy)[ends]
    recall = tp/y.sum()
    precision = tp/(ends+1)
    return float(np.sum(np.diff(np.r_[0.,recall])*precision))


def auc(y, s):
    y = np.asarray(y, dtype=bool)
    n1, n0 = int(y.sum()), int((~y).sum())
    if not n1 or not n0:
        return float('nan')
    return float((rankdata(s)[y].sum()-n1*(n1+1)/2)/(n1*n0))


def spearman(a, b):
    if len(a)<2 or np.std(a)==0 or np.std(b)==0:
        return float('nan')
    return float(np.corrcoef(rankdata(a),rankdata(b))[0,1])


def match_tables(a, d):
    design = pd.read_csv(DOCS/'matched_compound_design.tsv', sep='\t')
    rows = []
    for _, item in design[design.status=='MATCHED'].iterrows():
        x = a[(a.SC_cat==item.SC_cat) & np.isclose(a.dose_M,item.primary_dose_M,rtol=1e-8,atol=0)]
        z = d[(d.SC_cat==item.SC_cat) & np.isclose(d.dose_M,item.followup_dose_M,rtol=1e-8,atol=0)]
        if len(x)!=1 or len(z)!=1:
            raise ValueError('Locked compound/dose missing or duplicated')
        row = {'SC_cat':item.SC_cat,'compound_name':item.followup_name,
               'primary_dose_M':item.primary_dose_M,'followup_dose_M':item.followup_dose_M}
        for prefix, t in [('screen',x.iloc[0]),('followup',z.iloc[0])]:
            row.update({prefix+'_'+k:v for k,v in t.items() if k not in ['SC_cat','dose_M','compound_name']})
        row['eligible'] = bool(row['screen_eligible'] and row['followup_eligible'])
        rows.append(row)
    return pd.DataFrame(rows)


def benchmark(table, retention=0.8):
    t = table[table.eligible].copy()
    y = t[f'followup_joint_response_{retention}'].to_numpy(bool)
    a, m = t.screen_alpha_score.to_numpy(), t.screen_multivariate_score.to_numpy()
    ap_a, ap_m = average_precision(y,a), average_precision(y,m)
    loo = []
    for i in range(len(t)):
        keep = np.arange(len(t)) != i
        loo.append(average_precision(y[keep],m[keep])-average_precision(y[keep],a[keep]))
    # An undefined AP comparison is missing, never a measured zero gain.
    loof = float(np.mean(np.array(loo)>0)) if len(loo) and np.isfinite(loo).all() else float('nan')
    evaluable = len(t)>=CONFIG['minimum_matched_compounds'] and min(int(y.sum()),int((~y).sum()))>=CONFIG['minimum_positive_and_negative_compounds']
    gain = ap_m-ap_a
    status = 'BENCHMARK_NOT_EVALUABLE'
    if evaluable:
        status = 'WITHIN_STUDY_PREDICTIVE_GAIN' if gain>=CONFIG['minimum_average_precision_gain'] and loof>=CONFIG['minimum_leave_one_compound_out_positive_fraction'] else 'NO_DEMONSTRATED_PREDICTIVE_GAIN'
    result = {'n_compounds':len(t),'positive':int(y.sum()),'negative':int((~y).sum()),
              'retention':retention,'alpha_AP':ap_a,'multivariate_AP':ap_m,'AP_gain':gain,
              'alpha_AUC':auc(y,a),'multivariate_AUC':auc(y,m),
              'alpha_repeatability_spearman':spearman(a,t.followup_alpha_score),
              'multivariate_repeatability_spearman':spearman(m,t.followup_multivariate_score),
              'leave_one_compound_out_positive_gain_fraction':loof,'decision':status}
    return result, pd.DataFrame({'SC_cat':t.SC_cat,'compound_name':t.compound_name,'AP_gain_after_removal':loo})


def biochemical(out, followup):
    records, issues, vectors = [], [], []
    for assay, filename in [('ROS','ROS_all_compounds.xlsx'),('JC1','JC1_all_compounds.xlsx')]:
        wb = openpyxl.load_workbook(CACHE/'Fig4/Raw data'/filename, read_only=True,data_only=False)
        for ws in wb:
            data = list(ws.iter_rows(values_only=True))
            if not data:
                continue
            header = data[0]
            for ci, col in enumerate(header):
                if col is None:
                    continue
                col = str(col).strip()
                normalized = col.replace('μ','u').replace('µ','u').replace('Μ','M')
                match = re.fullmatch(r'(\d+(?:\.\d+)?)\s*(uM|nM)',normalized)
                dose = float(match[1]) * (1e-6 if match[2]=='uM' else 1e-9) if match else None
                kind = 'treatment' if dose is not None else 'control' if ('GC' in col or 'SNCA' in col) else 'unrecognized'
                if kind=='unrecognized':
                    issues.append({'assay':assay,'sheet':ws.title,'location':'header','raw':col,'reason':'unrecognized_column'})
                    continue
                vector = []
                for ri, row in enumerate(data[1:],start=2):
                    v = row[ci] if ci<len(row) else None
                    if v is None:
                        continue
                    raw = str(v).strip()
                    vector.append(raw)
                    flagged = '*' in raw
                    parsed = raw.replace('*','').strip()
                    if not re.fullmatch(r'[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?',parsed):
                        issues.append({'assay':assay,'sheet':ws.title,'location':f'{openpyxl.utils.get_column_letter(ci+1)}{ri}',
                                       'raw':raw,'reason':'malformed_numeric_not_repaired'})
                        continue
                    number = float(parsed)
                    if not math.isfinite(number):
                        continue
                    records.append({'assay':assay,'sheet':ws.title,'source_row':ri,'source_column':ci+1,
                                    'condition':col,'kind':kind,'dose_M':dose,'value':number,'flagged':flagged})
                if kind=='control' and vector:
                    vectors.append({'assay':assay,'sheet':ws.title,'condition':col,'n_entries':len(vector),
                                    'vector_sha256':hashlib.sha256(json.dumps(vector).encode()).hexdigest()})
        wb.close()
    long = pd.DataFrame(records)
    summaries = []
    for include in [False,True]:
        temp = long if include else long[~long.flagged]
        for (assay,sheet), frame in temp.groupby(['assay','sheet'],sort=True):
            ctrl = frame[frame.condition=='SNCA Trpl.'].value.median()
            corr = frame[frame.condition=='GC'].value.median()
            for (condition,kind), g in frame.groupby(['condition','kind'],sort=True):
                median = g.value.median()
                summaries.append({'assay':assay,'sheet':sheet,'condition':condition,'kind':kind,
                                  'dose_M':g.dose_M.iloc[0],'include_flagged':include,'n_values':len(g),
                                  'median':median,'mutant_control_median':ctrl,'corrected_control_median':corr,
                                  'ratio_to_mutant':median/ctrl if ctrl>0 else np.nan,
                                  'control_gap_fraction':(median-ctrl)/(corr-ctrl) if abs(corr-ctrl)>1e-6 else np.nan})
    summary = pd.DataFrame(summaries)
    names = followup[['SC_cat','compound_name']].drop_duplicates()
    aliases = {'SLK2001':'SKL2001','Tyrphostatin 9':'Tyrphostin 9',
               'Chloroquine':'Chloroquine Phosphate','Sertraline':'Sertraline HCl'}
    links = []
    for sheet in sorted(long[long.kind=='treatment'].sheet.unique()):
        target = aliases.get(sheet,sheet)
        found = names[names.compound_name==target]
        link = {'biochemical_sheet':sheet,'imaging_name':target,'join':'inferred_spelling_or_salt' if sheet in aliases else 'exact_name'}
        if len(found)!=1:
            link.update(join='UNMATCHED',SC_cat=None)
            links.append(link)
            continue
        cat = found.SC_cat.iloc[0]
        fg = followup[followup.SC_cat==cat]
        desired = CONFIG['biochemical_imaging_comparison_dose_M']
        nearest = fg.iloc[np.argmin(abs(fg.dose_M.to_numpy()-desired))]
        if abs(nearest.dose_M-desired)/desired>CONFIG['maximum_relative_dose_difference']:
            link.update(join='DOSE_MISMATCH',SC_cat=cat)
            links.append(link)
            continue
        link.update(SC_cat=cat,imaging_dose_M=nearest.dose_M,
                    alpha_score=nearest.alpha_score,multivariate_score=nearest.multivariate_score,
                    joint_response_0_8=nearest['joint_response_0.8'],
                    joint_fraction_0_8=nearest['joint_fraction_0.8'],
                    living_nuclei_ratio=nearest[VALUES[4]+'_ratio'])
        for assay in ['ROS','JC1']:
            ss = summary[(summary.sheet==sheet)&(summary.assay==assay)&(~summary.include_flagged)
                         &np.isclose(summary.dose_M,desired,atol=0,rtol=1e-8)]
            if len(ss)==1:
                link[assay+'_ratio_to_mutant'] = ss.ratio_to_mutant.iloc[0]
                link[assay+'_n_values'] = int(ss.n_values.iloc[0])
        links.append(link)
    vector_df = pd.DataFrame(vectors)
    vector_df['matching_sheet_count'] = vector_df.groupby(['assay','condition','vector_sha256']).sheet.transform('size')
    write_table(long,out/'biochemical_values.tsv.gz')
    write_table(summary,out/'biochemical_summary.tsv')
    write_table(vector_df,out/'biochemical_control_duplication.tsv')
    write_table(pd.DataFrame(issues),out/'biochemical_parse_issues.tsv')
    write_table(pd.DataFrame(links),out/'cross_assay_descriptive.tsv')
    return {'parsed_values':len(long),'flagged_values':int(long.flagged.sum()),'parse_issues':issues,
            'sheets':long.groupby('assay').sheet.nunique().to_dict(),
            'cross_assay_exact_names':sum(x['join']=='exact_name' for x in links),
            'cross_assay_inferred_names':sum(x['join']=='inferred_spelling_or_salt' for x in links),
            'biological_replicates':'unresolved; no inferential tests performed'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'results/v2/functional_rescue')
    args = parser.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    design_a,design_d = load_design()
    control = pd.read_csv(CACHE/'Fig1/Notebooks/Exp67_4plates_controls.csv')
    raw_a = pd.read_csv(CACHE/(MAIN+'all_raw_data.csv'),usecols=['Plate','Well']+VALUES)
    raw_d = pd.read_csv(CACHE/'Fig3/Notebooks/Exp72_DRC.csv',usecols=['Plate','Well']+VALUES)
    for frame in [raw_a,raw_d]:
        frame['Well'] = frame.Well.map(well)
    a = design_a.merge(raw_a,on=['Plate','Well'],validate='one_to_one')
    d = design_d.merge(raw_d,on=['Plate','Well'],validate='one_to_one')
    control, audit_c = normalize(control,'Mut - DMSO','WT - DMSO')
    a, audit_a = normalize(a,'Mut - DMSO','WT - DMSO')
    d, audit_d = normalize(d,'Mut;DMSO','WT;DMSO')
    model = fit_model(control,CONFIG['covariance_shrinkage'])
    qc = []
    for plate in sorted(control.Plate.unique()):
        fit = fit_model(control[control.Plate!=plate],CONFIG['covariance_shrinkage'])
        held = score(control[(control.Plate==plate)&control.tags.isin(['Mut - DMSO','WT - DMSO'])],fit)
        held = held[held.valid]
        y = held.tags=='WT - DMSO'
        qc.append({'held_out_plate':plate,'wells':len(held),'alpha_AUC':auc(y,held.alpha_score),
                   'multivariate_AUC':auc(y,held.multivariate_score)})
    sa,sd,sc = score(a,model),score(d,model),score(control,model)
    aa,dd = aggregate(sa),aggregate(sd)
    matched = match_tables(aa,dd)
    primary,loo = benchmark(matched)
    sensitivities = [dict(analysis='primary',**primary)]
    for t in CONFIG['retention_sensitivities']:
        result,_ = benchmark(matched,t)
        sensitivities.append(dict(analysis=f'retention_{t}',**result))
    for s in CONFIG['covariance_shrinkage_sensitivities']:
        alt = fit_model(control,s)
        table = match_tables(aggregate(score(a,alt)),aggregate(score(d,alt)))
        result,_ = benchmark(table)
        sensitivities.append(dict(analysis=f'shrinkage_{s}',**result))
    result,_ = benchmark(matched[matched.SC_cat!='S5232'])
    sensitivities.append(dict(analysis='exclude_Alectinib',**result))
    salts = matched[matched.SC_cat.isin(['S2128','S2167'])]
    if len(salts)==2:
        row = salts.iloc[0].copy()
        row['SC_cat'] = 'BAZEDOXIFENE_COLLAPSED'
        row['compound_name'] = 'Bazedoxifene salts collapsed'
        for field in ['screen_alpha_score','screen_multivariate_score','followup_alpha_score','followup_multivariate_score']:
            row[field] = salts[field].median()
        row['eligible'] = bool(salts.eligible.all())
        for t in THRESHOLDS:
            row[f'followup_joint_response_{t}'] = bool(salts[f'followup_joint_response_{t}'].all())
        reduced = pd.concat([matched[~matched.SC_cat.isin(['S2128','S2167'])],pd.DataFrame([row])],ignore_index=True)
        result,_ = benchmark(reduced)
        sensitivities.append(dict(analysis='collapse_Bazedoxifene_salts',**result))
    write_table(pd.concat([audit_c.assign(dataset='control'),audit_a.assign(dataset='screen'),audit_d.assign(dataset='followup')]),args.output/'normalization_audit.tsv')
    write_table(pd.DataFrame(qc),args.output/'control_plate_validation.tsv')
    for name,frame in [('screen_compounds',aa),('followup_doses',dd),('matched_benchmark',matched),('leave_one_compound_out',loo),('sensitivity_summary',pd.DataFrame(sensitivities))]:
        write_table(frame,args.output/(name+'.tsv'))
    for name,frame in [('screen_wells',sa),('followup_wells',sd),('control_wells',sc)]:
        write_table(frame,args.output/(name+'.tsv.gz'))
    bio = biochemical(args.output,dd)
    report = {'plan_commit':'61e73fd','normalization_amendment_commit':'57cdfae',
              'normalization_amendment':'docs/v2/functional_rescue/normalization_amendment.md',
              'primary':primary,'model':model,
              'biochemical_audit':bio,'screen_compounds':len(aa),'followup_compound_doses':len(dd),
              'matched_compounds':len(matched),'matched_eligible':int(matched.eligible.sum()),
              'screen_alpha_response':int(aa.alpha_response.sum()),
              'screen_joint_response':int(aa['joint_response_0.8'].sum()),
              'matched_followup_alpha_response':int(matched.followup_alpha_response.sum()),
              'matched_followup_joint_response':int(matched['followup_joint_response_0.8'].sum()),
              'invalid_wells':{'control':int((~control.valid).sum()),'screen':int((~a.valid).sum()),'followup':int((~d.valid).sum())},
              'interpretation':'Finite selected-panel prediction and descriptive assay evidence; no new molecular mechanism or independent biological confirmation.'}
    def clean(obj):
        if isinstance(obj,dict):return {k:clean(v) for k,v in obj.items()}
        if isinstance(obj,list):return [clean(x) for x in obj]
        if isinstance(obj,float) and not math.isfinite(obj):return None
        if isinstance(obj,np.generic):return clean(obj.item())
        return obj
    (args.output/'analysis_summary.json').write_text(json.dumps(clean(report),indent=2,allow_nan=False)+'\n')
    print(json.dumps(clean({k:v for k,v in report.items() if k not in ['model','biochemical_audit']}),indent=2))


if __name__=='__main__':
    main()
