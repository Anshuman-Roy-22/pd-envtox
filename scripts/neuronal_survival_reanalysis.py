#!/usr/bin/env python3
"""Reproduce the fixed Tian 2021 neuronal-survival summary-data reanalysis."""
import argparse
import gzip
import hashlib
import importlib.metadata
import io
import json
from pathlib import Path
import platform
import sys
import urllib.request

import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import rankdata

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs/v2/perturbation_feasibility"
URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41593-021-00862-0/MediaObjects/41593_2021_862_MOESM3_ESM.xlsx"
SHA256 = "82785da6b6fd4f9732a1e6bd701ab7fc2aa4d12f0d33fddc676054e37841a464"
DRAWS = 100000
SHARED = ["POMP", "PSMA4", "PSMB4", "PSMC5", "PSMD1", "PSMD2", "PSMD4", "PSMG1"]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_tsv(frame, path):
    frame.to_csv(path, sep="\t", index=False, float_format="%.17g", na_rep="NA", lineterminator="\n")


def table(workbook, mode, ao):
    sheet = workbook[f"Survival_{ao}_{mode}"]
    header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=3, values_only=True))
    if header != ("Gene", "TSS", "Phenotype"):
        raise ValueError(f"Unexpected header in {sheet.title}: {header}")
    records = []
    for gene, tss, phenotype in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        if gene is None:
            continue
        gene = str(gene).strip()
        if not gene or gene.lower().startswith(("negative", "non-target", "nontarget", "control")):
            continue
        try:
            phenotype = float(phenotype)
        except (TypeError, ValueError):
            continue
        if not np.isfinite(phenotype):
            continue
        records.append((gene, str(tss), phenotype))
    raw = pd.DataFrame(records, columns=["gene", "tss", "phenotype"])
    return raw.groupby("gene", sort=True).agg(
        phenotype=("phenotype", "mean"),
        source_rows=("phenotype", "size"),
        unique_tss=("tss", "nunique"),
    ).add_suffix("_" + ao)


def mode_data(workbook, mode):
    frame = table(workbook, mode, "plusAO").join(table(workbook, mode, "noAO"), how="inner")
    for ao in ["plusAO", "noAO"]:
        frame["percentile_" + ao] = (rankdata(frame["phenotype_" + ao], method="average") - 0.5) / len(frame)
    frame["delta_rank"] = frame.percentile_noAO - frame.percentile_plusAO
    frame["mode"] = mode
    return frame


def sample_sums(rng, values, count):
    """Uniform independent draws of subsets, without replacement within each draw.

    Rejection of repeated indices conditions iid uniform draws on distinctness,
    preserving the uniform distribution on ordered samples without replacement.
    """
    indices = rng.integers(len(values), size=(DRAWS, count))
    if count > 1:
        repeats = np.any(np.diff(np.sort(indices, axis=1), axis=1) == 0, axis=1)
        while repeats.any():
            indices[repeats] = rng.integers(len(values), size=(int(repeats.sum()), count))
            repeats = np.any(np.diff(np.sort(indices, axis=1), axis=1) == 0, axis=1)
    return values[indices].sum(axis=1)


def compare(frame, targets, pathway, name, bins, seed, out):
    frame = frame.copy()
    frame["bin"] = np.minimum(bins - 1, np.floor(bins * frame.percentile_plusAO).astype(int))
    frame["multirow"] = (frame.source_rows_plusAO > 1).astype(int)
    selected = frame.loc[targets].copy()
    background = frame.loc[~frame.index.isin(pathway)]
    rng = np.random.Generator(np.random.PCG64(seed))
    null = np.zeros(DRAWS)
    groups = []
    expectations = {}
    for (bin_id, multirow), group in selected.groupby(["bin", "multirow"], sort=True):
        controls = background.loc[(background.bin == bin_id) & (background.multirow == multirow)]
        k, n = len(group), len(controls)
        if n < max(20, 5 * k):
            return {"comparison": name, "status": "NOT_EVALUABLE", "reason": f"Insufficient matched controls in bin {bin_id}, multirow {multirow}: {n}"}
        values = controls.delta_rank.to_numpy()
        null += sample_sums(rng, values, k)
        mean = float(values.mean())
        for gene in group.index:
            expectations[gene] = mean
        groups.append({"comparison": name, "baseline_bins": bins, "bin": int(bin_id),
                       "multirow": int(multirow), "targets": ";".join(group.index),
                       "target_n": k, "control_n": n, "control_mean_delta": mean})
    null /= len(targets)
    observed = float(selected.delta_rank.mean())
    p_low = float((1 + (null <= observed).sum()) / (DRAWS + 1))
    p_high = float((1 + (null >= observed).sum()) / (DRAWS + 1))
    exact_expectation = float(np.mean(list(expectations.values())))
    quantiles = np.quantile(null, [0.025, 0.975])
    buffer = io.StringIO()
    np.savetxt(buffer, null, fmt="%.17g", header="mean_delta_rank", comments="")
    payload = buffer.getvalue().encode("utf-8")
    compressed = gzip.compress(payload, mtime=0)
    reference = out / f"{name}_reference.tsv.gz"
    temporary = reference.with_suffix(reference.suffix + ".tmp")
    temporary.write_bytes(compressed)
    if gzip.decompress(temporary.read_bytes()) != payload:
        raise IOError("Compressed reference verification failed")
    temporary.replace(reference)
    write_tsv(pd.DataFrame(groups), out / f"{name}_matching.tsv")
    selected["matched_expected_delta"] = [expectations[g] for g in selected.index]
    selected["excess_delta"] = selected.delta_rank - selected.matched_expected_delta
    write_tsv(selected.reset_index(), out / f"{name}_targets.tsv")
    if name == "primary_shared8_CRISPRi":
        loo = [{"omitted_gene": gene,
                "observed_mean_delta": float(selected.drop(gene).delta_rank.mean()),
                "matched_expectation": float(selected.drop(gene).matched_expected_delta.mean()),
                "excess_delta": float(selected.drop(gene).excess_delta.mean())}
               for gene in selected.index]
        write_tsv(pd.DataFrame(loo), out / "primary_leave_one_gene_out.tsv")
    return {"comparison": name, "status": "EVALUABLE", "target_n": len(targets),
            "universe_n": len(frame), "background_n": len(background),
            "baseline_bins": bins, "seed": seed, "draws": DRAWS,
            "observed_mean_delta": observed, "reference_mean": float(null.mean()),
            "exact_reference_expectation": exact_expectation,
            "observed_minus_reference_mean": observed - float(null.mean()),
            "observed_minus_exact_expectation": observed - exact_expectation,
            "reference_2.5pct": float(quantiles[0]), "reference_97.5pct": float(quantiles[1]),
            "p_lower": p_low, "p_upper": p_high, "p_two_sided": min(1.0, 2 * min(p_low, p_high))}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "data_raw/neuronal_perturbation/tian2021_s1.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "results/v2/neuronal_survival")
    parser.add_argument("--fetch", action="store_true", help="Download the locked public input if absent")
    args = parser.parse_args()
    if args.fetch and not args.input.exists():
        args.input.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(URL, timeout=60) as response:
            data = response.read(20000001)
        if hashlib.sha256(data).hexdigest() != SHA256:
            raise ValueError("Downloaded input checksum mismatch; input not saved")
        args.input.write_bytes(data)
    if digest(args.input) != SHA256:
        raise ValueError("Input workbook checksum mismatch")
    args.output.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((DOCS / "candidate_manifest.json").read_text())
    pathway = set(manifest["pathway_genes"])
    if len(pathway) != 52 or SHARED != manifest["prior_shared_leading_edge_genes"]:
        raise ValueError("Candidate manifest changed")
    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    frames = {mode: mode_data(workbook, mode) for mode in ["CRISPRi", "CRISPRa"]}
    workbook.close()
    for mode, frame in frames.items():
        write_tsv(frame.reset_index(), args.output / f"{mode}_gene_level.tsv")
    comparisons = [
        ("primary_shared8_CRISPRi", "CRISPRi", SHARED, 20, 20260914),
        ("secondary_pathway_CRISPRi", "CRISPRi", sorted(pathway & set(frames["CRISPRi"].index)), 20, 20260915),
        ("secondary_shared8_CRISPRa", "CRISPRa", SHARED, 20, 20260916),
        ("sensitivity_10bins_CRISPRi", "CRISPRi", SHARED, 10, 20260917),
        ("sensitivity_40bins_CRISPRi", "CRISPRi", SHARED, 40, 20260918),
    ]
    summaries = []
    for name, mode, targets, bins, seed in comparisons:
        if name != "secondary_pathway_CRISPRi" and not set(SHARED) <= set(frames[mode].index):
            result = {"comparison": name, "status": "NOT_EVALUABLE", "reason": "Missing primary target"}
        else:
            result = compare(frames[mode], targets, pathway, name, bins, seed, args.output)
        summaries.append(result)
        print(json.dumps(result), flush=True)
    write_tsv(pd.DataFrame(summaries), args.output / "comparison_summary.tsv")
    primary = summaries[0]
    label = "NOT_EVALUABLE"
    if primary["status"] == "EVALUABLE":
        label = "CONDITIONAL_RANK_SHIFT" if primary["p_two_sided"] < 0.05 else "NO_CONDITIONAL_RANK_SHIFT"
    provenance = {"input_url": URL, "input_sha256": SHA256,
                  "analysis_plan_commit": "6d493a3", "primary_label": label,
                  "interpretation_scope": "Conditional gene-set reference test on published summary phenotypes; not independent PD replication or an experimental interaction test.",
                  "candidate_manifest_sha256": digest(DOCS / "candidate_manifest.json"),
                  "analysis_plan_sha256": digest(ROOT / "docs/v2/neuronal_survival_frozen_plan.md"),
                  "script_sha256": digest(Path(__file__)),
                  "environment": {"python": platform.python_version(), **{p: importlib.metadata.version(p) for p in ["numpy", "pandas", "scipy", "openpyxl"]}}}
    (args.output / "provenance.json").write_text(json.dumps(provenance, indent=2) + "\n")
    for name, _, _, _, _ in comparisons:
        reference = args.output / f"{name}_reference.tsv.gz"
        if reference.exists():
            data = gzip.decompress(reference.read_bytes())
            if len(data.splitlines()) != DRAWS + 1:
                raise IOError(f"Incomplete reference output: {reference.name}")
    print("PRIMARY_LABEL=" + label, flush=True)


if __name__ == "__main__":
    main()
